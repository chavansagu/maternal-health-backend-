"""
Client "User_template.xlsx" bulk upload  -  matching + dry run + save.

Input : the parsed file from client_template_reader.parse_client_workbook()
Output: a plan + preview rows telling, for every centre / ward / user in the file,
        whether it is  new  |  existing  |  invalid  (with the reason).

Rules (agreed with the client requirement)
  * District must be the uploader's own district. A Block that does not exist yet is CREATED (code = its name).
  * Centre is matched by NAME (or CODE) inside the district. Not found -> a new centre is created
    with an auto code (= its name, made unique). Found -> the centre is NOT edited.
  * Sub-centre wards: every ward in the sheet is looked up by NAME inside the block; missing ones are
    created. Only the listed wards are mapped to the sub-centre (never "all wards of the block").
  * A user is created for every valid user row; username / email must not already exist.

Design: load_catalog() does the few DB queries; analyze_client_groups() is pure Python (no DB) so it
can be unit-tested and the same plan is re-built on confirm (never trust the preview).
"""
import io
import re

from sqlalchemy import func
from sqlalchemy.exc import IntegrityError

import admin_bulk_upload_service as base
from auth import get_password_hash
from models import (
    District, Block, Ward, SubCentre, USGCentre, DeliveryPoint, PMSMACentre, User, UserRole,
    UserWardMapping, WardSubcentreMapping,
)


# reader type -> (engine entity name used by admin_bulk_upload_service, model, uses block)
TYPES = {
    "sub_centre":     {"entity": "sub-centres",     "model": SubCentre,     "role": "sub_centre",  "block_required": True},
    "pmsma":          {"entity": "pmsma-centres",   "model": PMSMACentre,   "role": "pmsma",       "block_required": True},
    "usg":            {"entity": "usg-centres",     "model": USGCentre,     "role": "usg_centre",  "block_required": True},
    "delivery_point": {"entity": "delivery-points", "model": DeliveryPoint, "role": "dp",          "block_required": False},
}
CODE_MAX = 50


def _n(v):
    return base._norm_name(v or "")


def _unique_code(taken, name):
    """Code = the name itself (this is how the existing data looks, e.g. block 'Agalpur' has code
    'Agalpur'). If that code is already used, add -2, -3 ... `taken` holds UPPER-case codes."""
    root = re.sub(r"\s+", " ", (name or "").strip())[:CODE_MAX] or "CODE"
    code, i = root, 2
    while code.upper() in taken:
        suffix = f"-{i}"
        code = root[: CODE_MAX - len(suffix)] + suffix
        i += 1
    taken.add(code.upper())
    return code


# ----------------------------------------------------------------------------
# DB -> catalog (only reads)
# ----------------------------------------------------------------------------
def load_catalog(db, user, parsed):
    """Reads everything the analysis needs. Few queries, no per-row DB access."""
    if not user.district_id:
        raise ValueError("Your account is not linked to a district")
    district = db.query(District).filter(District.id == user.district_id).first()
    if not district:
        raise ValueError("Your district was not found")

    blocks = db.query(Block).filter(Block.district_id == district.id, Block.is_active == True).all()  # noqa: E712
    block_ids = [b.id for b in blocks]

    wards = db.query(Ward).filter(Ward.block_id.in_(block_ids)).all() if block_ids else []

    centers, codes = {}, {}
    for t, cfg in TYPES.items():
        M = cfg["model"]
        codes[t] = {(c or "").upper() for (c,) in db.query(M.code).all()}   # code is unique GLOBALLY
        if t == "sub_centre":
            rows = db.query(M).filter(M.block_id.in_(block_ids)).all() if block_ids else []
        else:
            rows = db.query(M).filter(M.district_id == district.id).all()
        centers[t] = [{"id": r.id, "name": r.name, "code": r.code,
                       "block_id": getattr(r, "block_id", None), "is_active": bool(r.is_active)} for r in rows]

    usernames = {x["username"].lower() for g in parsed["groups"] for x in g["users"]}
    emails = {x["email"].lower() for g in parsed["groups"] for x in g["users"] if x["email"]}
    taken_users, taken_emails = set(), set()
    if usernames:
        taken_users = {u.lower() for (u,) in db.query(User.username).filter(func.lower(User.username).in_(usernames)).all()}
    if emails:
        taken_emails = {e.lower() for (e,) in db.query(User.email).filter(func.lower(User.email).in_(emails)).all() if e}

    return {
        "district": {"id": district.id, "name": district.name},
        "blocks": {_n(b.name): {"id": b.id, "name": b.name} for b in blocks},
        "wards": {(w.block_id, _n(w.name)): {"id": w.id, "name": w.name, "code": w.code, "is_active": bool(w.is_active)}
                  for w in wards},
        "ward_codes": _group_codes(wards),
        "block_codes": {(c or "").upper() for (c,) in db.query(Block.code).all()},   # block code is unique GLOBALLY
        "centers": centers,
        "codes": codes,
        "usernames": taken_users,
        "emails": taken_emails,
    }


def _group_codes(wards):
    out = {}
    for w in wards:
        out.setdefault(w.block_id, set()).add((w.code or "").upper())
    return out


# ----------------------------------------------------------------------------
# Pure analysis (no DB)
# ----------------------------------------------------------------------------
def _find_center(catalog, t, name, code_like, block_id, block_required):
    """Returns (match, error). Match by name or code inside the district."""
    key = _n(name)
    cands = [c for c in catalog["centers"][t] if _n(c["name"]) == key or (c["code"] or "").upper() == code_like.upper()]
    if not cands:
        return None, None
    if block_id is not None:
        same = [c for c in cands if c["block_id"] == block_id]
        if same:
            cands = same
        elif block_required or all(c["block_id"] is not None for c in cands):
            return None, f"'{name}' already exists under a different block"
    if len(cands) > 1:
        return None, f"'{name}' matches more than one existing record"
    return cands[0], None


def analyze_client_groups(parsed, catalog):
    """
    Returns {"plan": [group plans], "rows": [preview rows], "summary": {...}}.
    Each group plan:
      {type, entity, label, sheet, center, block_id, district_id, center_action: new|existing|invalid,
       center_id, code, wards:[{name, action: new|existing, id, code}], users:[{..., status, reason}],
       reason}
    """
    own = catalog["district"]
    taken_codes = {t: set(c) for t, c in catalog["codes"].items()}
    ward_codes = {b: set(c) for b, c in catalog["ward_codes"].items()}
    taken_block_codes = set(catalog.get("block_codes", set()))
    new_blocks = {}                      # norm block name -> {"key","name","code","district_id"}
    rows, plans = [], []

    summary = {"new_blocks": 0, "new_centres": 0, "existing_centres": 0, "new_wards": 0, "existing_wards": 0,
               "new_users": 0, "invalid": 0}

    # row-level problems found while reading the file
    for e in parsed["errors"]:
        rows.append({"row_number": e["row_number"], "sheet": e["sheet"], "kind": "user", "record_type": "User",
                     "name": "", "center": "", "block": "", "status": "invalid", "reason": e["reason"]})
        summary["invalid"] += 1

    for g in parsed["groups"]:
        cfg = TYPES[g["type"]]
        t = g["type"]
        plan = {"type": t, "entity": cfg["entity"], "role": cfg["role"], "label": g["label"], "sheet": g["sheet"],
                "center": g["center"], "block": g["block"], "district_id": own["id"], "block_id": None,
                "center_action": None, "center_id": None, "code": None, "wards": [], "users": [], "reason": None}
        first_row = g["row_numbers"][0]
        plan["first_row"] = first_row

        def bad(reason, plan=plan, g=g, first_row=first_row):
            plan["center_action"], plan["reason"] = "invalid", reason
            rows.append({"row_number": first_row, "sheet": g["sheet"], "kind": "center", "record_type": g["label"],
                         "name": g["center"], "center": g["center"], "block": g["block"] or "",
                         "status": "invalid", "reason": reason})
            summary["invalid"] += 1
            for u in g["users"]:
                plan["users"].append({**u, "status": "invalid", "reason": f"Skipped: {reason}"})
                rows.append({"row_number": u["row_number"], "sheet": g["sheet"], "kind": "user", "record_type": "User",
                             "name": u["full_name"], "username": u["username"], "center": g["center"],
                             "block": g["block"] or "", "status": "invalid", "reason": f"Skipped: {reason}"})
                summary["invalid"] += 1
            plans.append(plan)

        # ---- district ----
        if g["district"] and _n(g["district"]) != _n(own["name"]):
            bad(f"District '{g['district']}' is outside your district ('{own['name']}')")
            continue
        # ---- block ----
        blk = None
        if g["block"]:
            bkey = _n(g["block"])
            blk = catalog["blocks"].get(bkey)
            if blk:
                plan["block_id"] = blk["id"]
            else:
                # block not in DB -> it will be created (code = its name). id is a "new:..." placeholder
                nb = new_blocks.get(bkey)
                if not nb:
                    nb = {"key": f"new:{bkey}", "name": g["block"], "district_id": own["id"],
                          "code": _unique_code(taken_block_codes, g["block"]),
                          "sheet": g["sheet"], "row_number": first_row}
                    new_blocks[bkey] = nb
                plan["block_id"] = nb["key"]
        elif cfg["block_required"]:
            bad("Block is required")
            continue

        # ---- centre ----
        match, err = _find_center(catalog, t, g["center"], g["center"], plan["block_id"], cfg["block_required"])
        if err:
            bad(err)
            continue
        if match:
            if t == "delivery_point" and not match["is_active"] and g["users"]:
                bad(f"Delivery Point '{match['name']}' is inactive; users cannot be assigned to it")
                continue
            plan["center_action"], plan["center_id"], plan["code"] = "existing", match["id"], match["code"]
            if plan["block_id"] is None:
                plan["block_id"] = match["block_id"]       # DP sheet has no block: use the centre's own
            summary["existing_centres"] += 1
        else:
            plan["center_action"] = "new"
            plan["code"] = _unique_code(taken_codes[t], g["center"])
            summary["new_centres"] += 1

        # ---- wards (sub-centre only) ----
        for wname in g["wards"]:
            w = catalog["wards"].get((plan["block_id"], _n(wname)))
            if w:
                plan["wards"].append({"name": w["name"], "action": "existing", "id": w["id"], "code": w["code"]})
                summary["existing_wards"] += 1
            else:
                codes = ward_codes.setdefault(plan["block_id"], set())
                plan["wards"].append({"name": wname, "action": "new", "id": None,
                                      "code": _unique_code(codes, wname)})
                summary["new_wards"] += 1

        rows.append({"row_number": first_row, "sheet": g["sheet"], "kind": "center", "record_type": g["label"],
                     "name": g["center"], "center": g["center"], "block": g["block"] or "",
                     "status": plan["center_action"],
                     "reason": ("Already exists - not changed" if plan["center_action"] == "existing"
                                else f"Will be created (code: {plan['code']})"),
                     "wards_new": sum(1 for w in plan["wards"] if w["action"] == "new"),
                     "wards_existing": sum(1 for w in plan["wards"] if w["action"] == "existing")})

        # ---- users ----
        for u in g["users"]:
            reason = None
            if u["username"].lower() in catalog["usernames"]:
                reason = f"User Name '{u['username']}' already exists"
            elif u["email"] and u["email"].lower() in catalog["emails"]:
                reason = f"Email '{u['email']}' already exists"
            st = "invalid" if reason else "new"
            plan["users"].append({**u, "status": st, "reason": reason})
            rows.append({"row_number": u["row_number"], "sheet": g["sheet"], "kind": "user", "record_type": "User",
                         "name": u["full_name"], "username": u["username"], "center": g["center"],
                         "block": g["block"] or "", "status": st,
                         "reason": reason or ("Will be created" + (f" with {len(plan['wards'])} ward(s)" if plan["wards"] else ""))})
            summary["invalid" if reason else "new_users"] += 1
        plans.append(plan)

    used = {p["block_id"] for p in plans if isinstance(p["block_id"], str) and p["center_action"] != "invalid"}
    blocks_to_create = [nb for nb in new_blocks.values() if nb["key"] in used]
    for nb in blocks_to_create:
        rows.append({"row_number": nb["row_number"], "sheet": nb["sheet"], "kind": "block", "record_type": "Block",
                     "name": nb["name"], "center": "", "block": nb["name"], "status": "new",
                     "reason": f"New block will be created (code: {nb['code']}) - check the spelling"})
    summary["new_blocks"] = len(blocks_to_create)

    rows.sort(key=lambda r: (r["sheet"], r["row_number"], 0 if r["kind"] == "block" else 1))
    return {"plan": plans, "rows": rows, "summary": summary, "new_blocks": blocks_to_create}


# ----------------------------------------------------------------------------
# Save (re-uses the live engine for blocks / centres / wards, live rules for users)
# ----------------------------------------------------------------------------
_ROLE = {"sub_centre": UserRole.SUB_CENTRE, "pmsma": UserRole.PMSMA,
         "usg_centre": UserRole.USG_CENTRE, "dp": UserRole.DP}
_CENTER_FK = {"sub_centre": "sub_centre_id", "pmsma": "pmsma_centre_id",
              "usg": "usg_centre_id", "delivery_point": "dp_id"}


def execute_client_plan(db, user, analysis):
    """
    Saves the analysed plan. One SAVEPOINT per centre-group (centre + wards + mappings) and one per user,
    so one bad row never blocks the rest. The caller commits.
    Returns counts + failed rows. Passwords are never returned / logged.
    """
    out = {"blocks": 0, "centres": 0, "wards": 0, "users": 0, "ward_mappings": 0,
           "failed": 0, "failed_rows": [], "errors": []}

    def fail(sheet, row_number, name, reason, record_type=""):
        out["failed"] += 1
        out["failed_rows"].append({"row_number": row_number, "sheet": sheet, "record_type": record_type,
                                   "name": name, "code": "", "error_reason": reason})
        out["errors"].append(f"{sheet} row {row_number} | {name} | Reason: {reason}")

    # rows that were already invalid in the analysis (bad file rows, duplicate username ...)
    for r in analysis["rows"]:
        if r["status"] == "invalid":
            fail(r["sheet"], r["row_number"], r.get("username") or r.get("name") or r.get("center", ""),
                 r["reason"], r.get("record_type", ""))

    ward_cfg = base.get_config("wards")

    # new blocks first (own savepoint each); groups that need a failed block are skipped below
    block_ids, block_errors = {}, {}
    for nb in analysis.get("new_blocks", []):
        try:
            with db.begin_nested():
                bobj = base._create(db, "blocks", base.get_config("blocks"),
                                    {"name": nb["name"], "name_regional": None, "code": nb["code"],
                                     "district_id": nb["district_id"]}, user)
            block_ids[nb["key"]] = bobj.id
            out["blocks"] += 1
        except Exception as e:
            block_errors[nb["key"]] = str(e)[:150]

    for p in analysis["plan"]:
        if p["center_action"] == "invalid":
            continue
        if isinstance(p["block_id"], str):                 # placeholder of a block created above
            if p["block_id"] in block_errors:
                reason = f"Could not create block '{p['block']}': {block_errors[p['block_id']]}"
                fail(p["sheet"], p.get("first_row") or 0, p["center"], reason, p["label"])
                for u in p["users"]:
                    if u["status"] == "new":
                        fail(p["sheet"], u["row_number"], u["username"], "Skipped: block could not be created", "User")
                continue
            p["block_id"] = block_ids[p["block_id"]]
        done = {"centres": 0, "wards": 0, "ward_mappings": 0}
        ward_ids, center_id = [], p["center_id"]
        try:
            with db.begin_nested():                       # centre + wards + mappings
                if p["center_action"] == "new":
                    d = {"name": p["center"], "code": p["code"], "district_id": p["district_id"],
                         "block_id": p["block_id"], "_block_ids": [p["block_id"]] if p["block_id"] else []}
                    obj = base._create(db, p["entity"], base.get_config(p["entity"]), d, user, map_all_wards=False)
                    center_id = obj.id
                    done["centres"] += 1

                for w in p["wards"]:
                    if w["action"] == "new":
                        wobj = base._create(db, "wards", ward_cfg,
                                            {"name": w["name"], "name_regional": None, "code": w["code"],
                                             "block_id": p["block_id"]}, user)
                        ward_ids.append(wobj.id)
                        done["wards"] += 1
                    else:
                        ward_ids.append(w["id"])

                if p["type"] == "sub_centre" and ward_ids:
                    have = {wid for (wid,) in db.query(WardSubcentreMapping.ward_id).filter(
                        WardSubcentreMapping.sub_centre_id == center_id,
                        WardSubcentreMapping.ward_id.in_(ward_ids)).all()}
                    for wid in ward_ids:
                        if wid not in have:
                            db.add(WardSubcentreMapping(ward_id=wid, sub_centre_id=center_id))
                            done["ward_mappings"] += 1
                    db.flush()
        except Exception as e:      # whole group rolled back
            msg = str(e)[:150]
            fail(p["sheet"], p.get("first_row") or 0, p["center"], f"Could not save centre/wards: {msg}", p["label"])
            for u in p["users"]:
                if u["status"] == "new":
                    fail(p["sheet"], u["row_number"], u["username"], "Skipped: centre/wards could not be saved", "User")
            continue
        for k in done:
            out[k] += done[k]

        for u in p["users"]:
            if u["status"] != "new":
                continue
            try:
                with db.begin_nested():
                    fk = {_CENTER_FK[p["type"]]: center_id}
                    new_user = User(
                        username=u["username"], email=u["email"], password_hash=get_password_hash(u["password"]),
                        role=_ROLE[p["role"]], full_name=u["full_name"], mobile_number=u["mobile"],
                        district_id=p["district_id"], block_id=p["block_id"], created_by=user.id, **fk)
                    db.add(new_user)
                    db.flush()
                    if p["type"] == "sub_centre":
                        for wid in ward_ids:
                            db.add(UserWardMapping(user_id=new_user.id, ward_id=wid))
                        db.flush()
                out["users"] += 1
            except IntegrityError:
                fail(p["sheet"], u["row_number"], u["username"], "Database rejected the user (User Name / Email already exists)", "User")
            except Exception as e:
                fail(p["sheet"], u["row_number"], u["username"], str(e)[:150], "User")
    out["failed_rows"].sort(key=lambda r: (r["sheet"], r["row_number"]))
    return out


# ----------------------------------------------------------------------------
# Template (xlsx) - same 4 sheets as the client's file. Data sheets have headers ONLY so a sample row
# can never be uploaded by mistake; examples live on the Instructions sheet.
# ----------------------------------------------------------------------------
_TEMPLATE_SHEETS = [
    ("Subcenter", ["District", "Block", "Sub center", "Full Name", "Mobile no", "Email ID", "User Name", "Password",
                   "Assigned ward/village"]),
    ("PMSMA", ["District", "Block", "PMSMA Center", "Full Name", "Mobile no", "Email ID", "User Name", "Password"]),
    ("Usg User", ["District", "Block", "USG Center", "Full Name", "Mobile no", "Email ID", "User Name", "Password"]),
    ("DP User", ["District", "Delivery Point", "Full Name", "Mobile no", "Email ID", "User Name", "Password"]),
]


def build_client_template_xlsx(district_name, block_names):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="FFF2CC")
    first = True
    for title, headers in _TEMPLATE_SHEETS:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append(headers)
        for i, h in enumerate(headers, start=1):
            c = ws.cell(1, i)
            c.font, c.fill = Font(bold=True), head_fill
            ws.column_dimensions[get_column_letter(i)].width = max(16, len(h) + 4)
            if h in ("Mobile no", "Password", "User Name"):
                for r in range(2, 502):
                    ws.cell(r, i).number_format = "@"          # keep as text (no 8.8E+10, no lost zeros)
            if h == "District" and district_name:
                ws.cell(2, i).value = district_name             # fixed district, written once on the first row
        ws.freeze_panes = "A2"

    ins = wb.create_sheet("Instructions")
    lines = [
        ("HOW TO FILL THIS FILE", True),
        ("1. Use the 4 sheets: Subcenter, PMSMA, Usg User, DP User. Fill only the sheets you need. Do NOT rename them.", False),
        ("2. Each row is one user. The centre (Sub center / PMSMA / USG / Delivery Point) is created if it does not exist yet.", False),
        ("3. District is already filled in (do not change it). Block and Centre can be written only on the first row; leave the cells below blank for more users of the same centre.", False),
        ("4. Subcenter sheet: list every ward/village of the sub-centre in 'Assigned ward/village', one per row (any number). "
         "Missing wards are created in that Block and mapped to the sub-centre; all users of the sub-centre get these wards.", False),
        ("5. Block: use an existing block name exactly (see the list below). If you write a NEW block name, that block is created - so check the spelling.", False),
        ("6. Mobile no: 10 digits. Email ID must be different for every user. User Name must be unique.", False),
        ("7. Password: at least 8 characters with an upper-case letter, a lower-case letter, a number and a special character "
         "(example: Sakma@1234).", False),
        ("8. An existing centre is NOT changed; only the new users (and wards) are added to it.", False),
        ("9. Upload the file; a review screen shows what will be created before anything is saved.", False),
        ("", False),
        ("EXAMPLE (Subcenter sheet)", True),
        ("District: Balangir | Block: Balangir | Sub center: Sakma | Full Name: Ravi Kumar | Mobile no: 9876543210 | "
         "Email ID: ravi@example.com | User Name: sakma_mphwf | Password: Sakma@1234 | Assigned ward/village: Bhainsa", False),
        ("", False),
        ("YOUR DISTRICT", True),
        (district_name or "", False),
        ("", False),
        ("EXISTING BLOCK NAMES (copy exact names)", True),
    ] + [(b, False) for b in block_names]
    for item in lines:
        text = item[0]
        bold = item[1] if len(item) > 1 else False
        ins.append([text])
        if bold:
            ins.cell(ins.max_row, 1).font = Font(bold=True)
    ins.column_dimensions["A"].width = 130
    for row in ins.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()